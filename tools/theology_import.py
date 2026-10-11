#!/usr/bin/env python3
# 신학(Theology) 파트 가져오기: Dropbox /.Claude_신학 의 정리 문서(docx) → data/theology.json
# 사용: python3 -I tools/theology_import.py <docx가 있는 폴더> tools/prayer_app.src.html data/theology.json
#  (두 번째 인자: 앱에 내장된 개역한글 성경 — 문서에 없는 구절의 대체 본문·장 길이 계산용)
#  · 표(구절 | 말씀 | 의미/적용) → 말씀 카드  · Heading → 소주제  · 장 단위 본문 → 본문 전체
#  · 본문은 문서에 실린 개역개정을 그대로 사용. 문서에 없는 구절(권면 매뉴얼의 '본문 근거')은
#    다른 문서들에서 모은 개역개정 절 색인에서 찾고, 그래도 없으면 개역한글로 대신하며 표시(alt=1)
import sys, re, json, os, datetime
import docx
from docx.table import Table
from docx.text.paragraph import Paragraph

SRC, BIBLE_JSON, OUT = sys.argv[1], sys.argv[2], sys.argv[3]
if BIBLE_JSON.endswith(".html"):  # 앱 소스(prayer_app.src.html)에 내장된 성경을 바로 사용
    import base64, gzip
    _m = re.search(r'<script id="bible" type="text/plain">([^<]+)</script>', open(BIBLE_JSON, encoding="utf-8").read())
    B = json.loads(gzip.decompress(base64.b64decode(_m.group(1))).decode("utf-8"))
else:
    B = json.load(open(BIBLE_JSON, encoding="utf-8"))
CODES, NAMES, KRV = B["codes"], B["names"], B["kr"]
ABBR = "창 출 레 민 신 수 삿 룻 삼상 삼하 왕상 왕하 대상 대하 스 느 에 욥 시 잠 전 아 사 렘 애 겔 단 호 욜 암 옵 욘 미 나 합 습 학 슥 말 마 막 눅 요 행 롬 고전 고후 갈 엡 빌 골 살전 살후 딤전 딤후 딛 몬 히 약 벧전 벧후 요일 요이 요삼 유 계".split()
BOOK = {}
for c, n, a in zip(CODES, NAMES, ABBR): BOOK[n] = c; BOOK[a] = c
BOOK["시편"] = "psa"; BOOK["요한일서"] = "1jn"
ONECH = {"oba", "phm", "2jn", "3jn", "jud"}
BK = "|".join(sorted(BOOK, key=len, reverse=True))
DASH = "[-–—~]"
REF_RE = re.compile(r"(?<![가-힣])(" + BK + r")\s?(\d{1,3})(?:(편|장)|:(\d{1,3}(?:\s?" + DASH + r"\s?(?:\d{1,3}:)?\d{1,3})?(?:\s?[,·]\s?\d{1,3}(?:\s?" + DASH + r"\s?\d{1,3})?)*))?")

def chlen(code, ch):
    try: return len(KRV[code][ch - 1])
    except Exception: return 0

def parse_ref(s):
    """'요 6:37,39' / '시편 22편' / '이사야 52:13–15' / '요 18:33-19:16' / '유 22' → (code,[(ch,v1,v2),...]) or None"""
    s = s.strip().replace("–", "-").replace("—", "-").replace("~", "-")
    m = re.match(r"^(" + BK + r")\s?(\d{1,3})(.*)$", s)
    if not m: return None
    code = BOOK[m.group(1)]; ch = int(m.group(2)); rest = m.group(3).strip()
    if rest in ("편", "장", ""):
        if rest == "" and code in ONECH: return (code, [(1, ch, ch)])
        if rest == "" and code not in ONECH: return (code, [(ch, 1, chlen(code, ch))])
        return (code, [(ch, 1, chlen(code, ch))])
    if not rest.startswith(":"): return None
    segs = []
    for part in re.split(r"\s?[,·]\s?", rest[1:]):
        mm = re.match(r"^(\d+)(?:\s?-\s?(?:(\d+):)?(\d+))?", part)
        if not mm: break
        v1 = int(mm.group(1))
        if mm.group(2):  # 장을 넘어가는 범위
            ch2, v2 = int(mm.group(2)), int(mm.group(3))
            segs.append((ch, v1, chlen(code, ch)))
            for c2 in range(ch + 1, ch2): segs.append((c2, 1, chlen(code, c2)))
            segs.append((ch2, 1, v2)); ch = ch2
        else:
            segs.append((ch, v1, int(mm.group(3)) if mm.group(3) else v1))
    return (code, segs) if segs else None

# ---------- 개역개정 절 색인 (문서에서 수집) ----------
KRR = {}  # (code,ch,v) -> text
def put(code, ch, v, t):
    t = t.strip()
    if t and (code, ch, v) not in KRR: KRR[(code, ch, v)] = t
def split_verses(text, segs):
    """'9하나님의…10사랑은…' 처럼 절 번호가 붙은 본문을 예상 절 번호 순서대로 나눔"""
    exp = [(ch, v) for ch, a, b in segs for v in range(a, b + 1)]
    if len(exp) == 1: return [(exp[0], re.sub(r"^\d+\s?", "", text) if re.match(r"^%d\D" % exp[0][1], text) else text)]
    out, pos = [], 0
    marks = []
    for (ch, v) in exp:
        m = re.compile(r"(?:^|(?<=[\s.!?」』\"'):;,]))%d(?=\S)" % v).search(text, pos)
        if not m: return None
        marks.append((m.start(), m.end(), (ch, v))); pos = m.end()
    for i, (s, e, cv) in enumerate(marks):
        nxt = marks[i + 1][0] if i + 1 < len(marks) else len(text)
        out.append((cv, text[e:nxt].strip()))
    return out

def lookup(code, segs):
    """개역개정 색인 우선, 없으면 개역한글. returns (verses[[ch,v,text]], alt)"""
    vs, alt = [], 0
    for ch, a, b in segs:
        for v in range(a, b + 1):
            t = KRR.get((code, ch, v))
            if t is None:
                try: t = KRV[code][ch - 1][v - 1]; alt = 1
                except Exception: continue
                t = re.sub(r"^<[^>]+>\s*", "", t)
            vs.append([ch, v, t])
    return vs, alt

def label(code, segs):
    nm = NAMES[CODES.index(code)]
    ch, a, b = segs[0]
    if len(segs) == 1 and a == 1 and b == chlen(code, ch): return f"{nm} {ch}{'편' if code == 'psa' else '장'}"
    return None

def blocks_of(doc):
    for el in doc.element.body.iterchildren():
        if el.tag.endswith("}p"): yield Paragraph(el, doc)
        elif el.tag.endswith("}tbl"): yield Table(el, doc)
def sty(p): return p.style.name if p.style is not None else ""
def cells(row): return [c.text.strip() for c in row.cells]
ROMAN = re.compile(r"^([ⅠⅡⅢⅣⅤⅥⅦⅧⅨⅩ])\.\s+(.+)$")

# ---------- 표형 문서: 하나님의 약속 / 기뻐하시는 것 ----------
def verse_table_doc(path, tid, title_override=None):
    d = docx.Document(path)
    T = {"id": tid, "src": os.path.basename(path), "intro": [], "parts": []}
    head, part, sub, pend = [], None, None, None
    def new_part(t, desc=""):
        nonlocal part, sub
        part = {"t": t, "d": desc, "subs": []}; T["parts"].append(part); sub = None
    for b in blocks_of(d):
        if isinstance(b, Paragraph):
            t = b.text.strip()
            if not t: continue
            st = sty(b)
            mr = ROMAN.match(t)
            if mr and not st.startswith("Heading"):
                new_part(f"{mr.group(1)}. {mr.group(2).strip()}"); pend = "desc"; continue
            if pend == "desc" and part is not None and not part["subs"] and not part["d"] and not st.startswith("Heading"):
                part["d"] = t; pend = None; continue
            if st.startswith("Heading") or re.match(r"^\d{1,2}\.\s+\S", t) and len(t) < 60:
                if part is None: new_part("")
                sub = {"t": t, "b": []}; part["subs"].append(sub); continue
            if part is None: head.append(t); continue
            if sub is None: (part.setdefault("x", [])).append(t); continue
            sub["b"].append({"k": "p", "t": t})
        else:
            rows = b.rows
            if len(rows) == 1 and len(b.columns) == 1:  # 장 구분 상자
                txt = [x.strip() for x in rows[0].cells[0].text.split("\n") if x.strip()]
                mr = ROMAN.match(txt[0]) if txt else None
                if mr:
                    new_part(f"{mr.group(1)}. {mr.group(2).strip()}", " ".join(txt[1:])); continue
                if txt and re.match(r"^제\s?\d장", txt[0]):
                    if not txt[0].startswith("제1장"): new_part(re.sub(r"\s+", " ", txt[0]), " ".join(txt[1:]))
                    continue
            hdr = cells(rows[0])
            if len(hdr) >= 3 and hdr[0] == "구절":
                if sub is None: continue
                for r in rows[1:]:
                    c = cells(r); ref, txt, note = c[0], c[1], c[2] if len(c) > 2 else ""
                    pr = parse_ref(ref); vv = None
                    if pr:
                        code, segs = pr; sp = split_verses(txt, segs)
                        if sp:
                            for (ch, v), t in sp: put(code, ch, v, t)
                            vv = [[ch, v, t] for (ch, v), t in sp]
                    sub["b"].append({"k": "v", "r": ref, "c": pr[0] if pr else None, "vs": vv, "x": None if vv else txt, "n": note})
            elif part is None:
                continue  # 한눈에 보기 표
            else:
                tgt = sub["b"] if sub else part.setdefault("xb", [])
                tgt.append({"k": "tb", "rows": [cells(r) for r in rows]})
    T["head"] = head
    return T

# ---------- 장 단위 본문: 십자가의 고난 ----------
def chapter_doc(path, tid):
    d = docx.Document(path)
    T = {"id": tid, "src": os.path.basename(path), "head": [], "parts": []}
    ps = list(blocks_of(d))
    part = sub = None; in_toc = True; cur = None; verse_mode = False; seen_toc = False
    for b in ps:
        if isinstance(b, Table):
            if part is not None and part["t"].startswith("부록"):
                (sub["b"] if sub else part.setdefault("xb", [])).append({"k": "tb", "rows": [cells(r) for r in b.rows]})
            continue
        t = b.text.strip(); st = sty(b)
        if not t: continue
        if st == "Heading 1":
            in_toc = (t == "차례"); seen_toc = True
            if in_toc: continue
            part = {"t": t, "d": "", "subs": []}; T["parts"].append(part); sub = None; continue
        if in_toc:
            if part is None and not T["parts"] and not seen_toc: T["head"].append(t)
            continue
        if part is None: continue
        if st == "Heading 2":
            pr = parse_ref(t.replace("편", "편").strip())
            sub = {"t": t, "n": "", "b": [], "ref": t}; part["subs"].append(sub); cur = pr; verse_mode = False; continue
        if st == "Heading 3":  # 부록의 소제목
            sub = {"t": t, "n": "", "b": []}; part["subs"].append(sub); cur = None; continue
        if sub is None:
            if not part["d"]: part["d"] = t
            continue  # 장 목록(차례 반복)은 건너뜀
        mv = re.match(r"^(\d{1,3})\s{1,3}(.+)$", t)
        if cur and mv:
            code, segs = cur; v = int(mv.group(1))
            ch = segs[0][0] if len(segs) == 1 else next((c for c, a, bb in segs if a <= v <= bb), segs[0][0])
            put(code, ch, v, mv.group(2))
            if not sub["b"] or sub["b"][-1]["k"] != "psg": sub["b"].append({"k": "psg", "code": code, "vs": []})
            sub["b"][-1]["vs"].append([ch, v, mv.group(2).strip()]); verse_mode = True; continue
        if cur and not sub["n"] and not verse_mode and not sub["b"]:
            sub["n"] = t; continue
        if cur:  # 소단락 제목 / 표제
            if not sub["b"] or sub["b"][-1]["k"] != "psg": sub["b"].append({"k": "psg", "code": cur[0], "vs": []})
            sub["b"][-1]["vs"].append([0, 0, t]); continue
        sub["b"].append({"k": "p", "t": t})
    return T

# ---------- 서술형 문서: 권면 매뉴얼 ----------
GROUND = re.compile(r"^본문 근거\s*[:：]\s*(.+)$")
def prose_doc(path, tid):
    d = docx.Document(path)
    T = {"id": tid, "src": os.path.basename(path), "head": [], "parts": []}
    part = sub = None; skip_toc = False
    for b in blocks_of(d):
        if isinstance(b, Table):
            if sub is not None: sub["b"].append({"k": "tb", "rows": [cells(r) for r in b.rows]})
            elif part is not None: part.setdefault("xb", []).append({"k": "tb", "rows": [cells(r) for r in b.rows]})
            continue
        t = b.text.strip(); st = sty(b)
        if not t: continue
        if st == "Heading 1":
            if t == "목차": skip_toc = True; continue
            skip_toc = False
            part = {"t": t, "d": "", "subs": []}; T["parts"].append(part); sub = None; continue
        if skip_toc: continue
        if part is None: T["head"].append(t); continue
        if st == "Heading 2":
            sub = {"t": t, "b": []}; part["subs"].append(sub); continue
        if st == "Heading 3":
            if sub is None: sub = {"t": t, "b": []}; part["subs"].append(sub); continue
            hb = {"k": "h", "t": t}
            mref = re.match(r"^(.+?)\s+[—-]\s+(.+)$", t)
            sub["b"].append(hb)
            if mref:
                pr = parse_ref(mref.group(1))
                if pr: sub["b"].append({"k": "v", "r": mref.group(1), "q": pr, "n": ""})
            continue
        tgt = sub["b"] if sub is not None else None
        if tgt is None:
            if not part["d"]: part["d"] = t
            else: part.setdefault("x", []).append(t)
            continue
        g = GROUND.match(t)
        if g:
            items = []
            for chunk in re.split(r"[,，]\s*(?![^()]*\))", g.group(1)):
                chunk = chunk.strip().rstrip(".")
                m = re.match(r"^(.+?)(?:\(([^)]*)\))?$", chunk)
                rr = m.group(1).strip(); lab = (m.group(2) or "").strip()
                pr = parse_ref(rr)
                if not pr and items:  # '마 6:24·33' 뒤의 같은 책 생략형 등
                    pr = None
                items.append({"r": rr, "l": lab, "q": pr})
            tgt.append({"k": "grd", "items": items}); continue
        tgt.append({"k": "li" if st == "List Paragraph" else "p", "t": t})
    return T

def resolve(T):
    """'q'(파싱된 참조)를 실제 본문으로 채움"""
    miss = []
    def fix(bl):
        if bl["k"] == "v" and "q" in bl:
            code, segs = bl.pop("q"); vs, alt = lookup(code, segs); bl["vs"] = vs; bl["c"] = code
            if alt: bl["alt"] = 1
        elif bl["k"] == "grd":
            for it in bl["items"]:
                q = it.pop("q", None)
                if q:
                    vs, alt = lookup(*q); it["vs"] = vs; it["c"] = q[0]
                    if alt: it["alt"] = 1
                else: miss.append(it["r"])
    for p in T["parts"]:
        for bl in p.get("xb", []): fix(bl)
        for s in p["subs"]:
            for bl in s["b"]: fix(bl)
    return miss

# 주제 목록: data/theology_manifest.json  [{id, t(제목), file(.Claude_신학 안 파일명), kind: table|prose|chapter}]
KIND = {"table": verse_table_doc, "prose": prose_doc, "chapter": chapter_doc}
MAN = json.load(open(os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "data", "theology_manifest.json"), encoding="utf-8"))
DOCS = [(m["id"], m["t"], m["file"], KIND[m["kind"]]) for m in MAN]
topics = []
for tid, title, fn, fnc in DOCS:
    p = os.path.join(SRC, fn)
    if not os.path.exists(p): print("missing", fn); continue
    T = fnc(p, tid); T["t"] = title; topics.append(T)
# 권면 문서는 다른 문서에서 모은 개역개정 색인을 쓰므로 맨 마지막에 해석
report = {}
for T in topics:
    report[T["id"]] = resolve(T)
for T in topics:  # 소주제 없이 본문만 있는 대단원 → 같은 이름의 소주제 하나로
    for P in T["parts"]:
        if not P["subs"] and (P.get("x") or P.get("xb") or P["d"]):
            P["subs"].append({"t": P["t"], "b": ([{"k": "p", "t": P["d"]}] if P["d"] else []) + [{"k": "p", "t": x} for x in P.pop("x", [])] + P.pop("xb", [])})
            P["d"] = ""
for T in topics:
    h = T.pop("head", [])
    T["sub"] = h[1] if len(h) > 1 else ""
    T["intro"] = [x for x in (h[2:] if len(h) > 2 else []) if "한눈에 보기" not in x and "이동합니다" not in x]
out = {"v": 1, "updated": datetime.date.today().isoformat(), "topics": topics}
json.dump(out, open(OUT, "w", encoding="utf-8"), ensure_ascii=False, separators=(",", ":"))
print("wrote", OUT, os.path.getsize(OUT) // 1024, "KB")
for T in topics:
    nv = sum(1 for p in T["parts"] for s in p["subs"] for b in s["b"] if b["k"] in ("v", "psg"))
    print(f"  {T['id']}: parts={len(T['parts'])} subs={sum(len(p['subs']) for p in T['parts'])} verse-blocks={nv} unresolved={report[T['id']][:8]}")
print("개역개정 색인 절 수:", len(KRR))
